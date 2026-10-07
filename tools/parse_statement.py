"""
Tool 1: parse_statement

Reads a bank / UPI statement (CSV, XLSX or PDF) and returns clean transactions:
    [{date, description, amount, type, balance}]
      date        -> "YYYY-MM-DD"
      amount      -> positive float
      type        -> "debit" (money out) or "credit" (money in)
      balance     -> float, or None if the file has no balance column

How it copes with different layouts (no per-bank code):
  1. Read the file into raw rows (list of lists).
  2. Find the real header row (skips junk lines above it).
  3. Map columns by keyword (Withdrawal/Debit/DR -> out, Deposit/Credit/CR -> in, ...).
  4. Normalise dates, amounts, Dr/Cr marks, rupee signs, commas.
"""
import csv
import logging
import re
from datetime import date, datetime
from pathlib import Path

log = logging.getLogger("statementiq.parse")


class ParseError(Exception):
    """Raised when a file cannot be understood. The agent shows this to the user."""


# ----------------------------------------------------------------- reading
def _read_csv(path: Path) -> list[list]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            with open(path, newline="", encoding=enc) as f:
                return [row for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise ParseError("Could not read the CSV file (unknown text encoding).")


def _read_xlsx(path: Path) -> list[list]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
        if any(any(str(c).strip() for c in r) for r in rows):
            return rows
    raise ParseError("The Excel file has no data.")


def _read_pdf(path: Path) -> list[list]:
    import pdfplumber
    rows, has_text = [], False
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            if (page.extract_text() or "").strip():
                has_text = True
            tables = page.extract_tables()
            if not tables:  # fallback for statements without ruled lines
                tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
            for t in tables:
                for r in t:
                    rows.append([(c or "").replace("\n", " ").strip() for c in r])
    if not has_text:
        raise ParseError("This looks like a scanned/image PDF, which is not supported. "
                         "Please upload the CSV or Excel version of your statement.")
    if not rows:
        raise ParseError("Could not find a transaction table in this PDF. Please upload a CSV instead.")
    return rows


# ----------------------------------------------------------------- cell helpers
DATE_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d %b %Y", "%d-%b-%Y", "%b %d, %Y", "%Y-%m-%d",
                "%d/%m/%y", "%d-%m-%y", "%d %B %Y", "%B %d, %Y", "%d.%m.%Y"]


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s:
        return None
    s = s.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2} \d", s) else s  # drop time part
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_amount(v):
    """Returns (value, direction) where direction is 'debit'/'credit'/None from a Dr/Cr mark."""
    if v is None or v == "":
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip()
    if not s:
        return None, None
    direction = None
    m = re.search(r"\(?\s*(dr|cr)\.?\s*\)?\s*$", s, re.I)
    if m:
        direction = "debit" if m.group(1).lower() == "dr" else "credit"
        s = s[:m.start()]
    neg = s.strip().startswith("-") or (s.strip().startswith("(") and s.strip().endswith(")"))
    digits = re.sub(r"[^0-9.]", "", s)
    if not digits or digits == ".":
        return None, direction
    val = float(digits)
    return (-val if neg else val), direction


def _clean(v) -> str:
    return " ".join(str(v).split())


# ----------------------------------------------------------------- layout detection
DESC_RE = r"narration|description|details|particulars|remarks"
AMT_RE = r"amount|amt|debit|credit|withdrawal|deposit|\bdr\b|\bcr\b"


def _find_header(rows):
    for i, row in enumerate(rows[:40]):
        cells = [_clean(c).lower() for c in row]
        if (any("date" in c for c in cells) and any(re.search(DESC_RE, c) for c in cells)
                and any(re.search(AMT_RE, c) for c in cells)):
            return i
    raise ParseError("Could not find the column headings (Date, Narration, Amount...). "
                     "Please upload a CSV exported from your bank app.")


def _map_columns(header):
    cells = [_clean(c).lower() for c in header]
    idx = {}
    for i, c in enumerate(cells):
        if "date" in c and "value" not in c and "date" not in idx:
            idx["date"] = i
        elif re.search(DESC_RE, c) and "desc" not in idx:
            idx["desc"] = i
        elif "balance" in c:
            idx["balance"] = i
        elif re.search(r"\b(withdrawal|debit|dr)\b", c):
            idx["debit"] = i
        elif re.search(r"\b(deposit|credit|cr)\b", c):
            idx["credit"] = i
        elif re.fullmatch(r"(type|dr/cr|cr/dr)", c):
            idx["type"] = i
        elif c == "status":
            idx["status"] = i
        elif re.search(r"amount|amt", c) and "amount" not in idx:
            idx["amount"] = i
    if "date" not in idx and any("date" in c for c in cells):  # only a "value date" exists
        idx["date"] = next(i for i, c in enumerate(cells) if "date" in c)
    if "date" not in idx or "desc" not in idx or not ({"debit", "credit", "amount"} & set(idx)):
        raise ParseError("Could not recognise the statement columns.")
    return idx


def _get(row, idx, key):
    i = idx.get(key)
    return row[i] if i is not None and i < len(row) else ""


# ----------------------------------------------------------------- main tool
def parse_statement(file_path: str) -> list[dict]:
    path = Path(file_path)
    if not path.exists():
        raise ParseError(f"File not found: {file_path}")
    ext = path.suffix.lower()
    if ext == ".csv":
        rows = _read_csv(path)
    elif ext in (".xlsx", ".xlsm"):
        rows = _read_xlsx(path)
    elif ext == ".pdf":
        rows = _read_pdf(path)
    else:
        raise ParseError(f"Unsupported file type '{ext}'. Please upload CSV, XLSX or PDF.")

    h = _find_header(rows)
    idx = _map_columns(rows[h])
    out, skipped = [], 0

    for row in rows[h + 1:]:
        if not any(_clean(c) for c in row):
            continue
        d = _parse_date(_get(row, idx, "date"))
        desc = _clean(_get(row, idx, "desc"))
        if d is None:
            # wrapped description line in a PDF: no date, no amounts -> glue to previous row
            has_amt = any(_clean(_get(row, idx, k)) for k in ("debit", "credit", "amount"))
            if out and desc and not has_amt and not _clean(_get(row, idx, "date")):
                out[-1]["description"] += " " + desc
            continue

        status = _clean(_get(row, idx, "status")).lower()
        if status and status not in ("success", "successful", "completed"):
            continue

        amount, kind = None, None
        dv, _ = _parse_amount(_get(row, idx, "debit"))
        cv, _ = _parse_amount(_get(row, idx, "credit"))
        if dv:
            amount, kind = dv, "debit"
        elif cv:
            amount, kind = cv, "credit"
        elif "amount" in idx:
            val, mark = _parse_amount(_get(row, idx, "amount"))
            if val is not None:
                amount = abs(val)
                kind = mark
                if kind is None and "type" in idx:
                    t = _clean(_get(row, idx, "type")).lower()
                    kind = "debit" if t.startswith(("d", "paid")) else "credit" if t.startswith(("c", "rec")) else None
                if kind is None:
                    kind = "debit" if val < 0 else "credit"
        if amount is None or kind is None:
            skipped += 1
            continue

        bal, bmark = _parse_amount(_get(row, idx, "balance"))
        if bal is not None and bmark == "debit":
            bal = -bal
        out.append({"date": d.isoformat(), "description": desc, "amount": round(abs(amount), 2),
                    "type": kind, "balance": None if bal is None else round(bal, 2)})

    if skipped:
        log.warning("Skipped %d rows that could not be read", skipped)
    if not out:
        raise ParseError("No transactions were found in this file.")
    return out


if __name__ == "__main__":
    import sys
    for t in parse_statement(sys.argv[1])[:10]:
        print(t)
